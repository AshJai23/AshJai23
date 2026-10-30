"""Synthetic edge cases only; never presented as actual retail evidence."""
import csv
import sys
import tempfile
import unittest
import zipfile
from datetime import datetime
from decimal import Decimal
from pathlib import Path

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from retail import FIELDS, build, normalise, workbook_to_csv
from metrics import query
import openpyxl

class IngestionTests(unittest.TestCase):
    def test_money_is_decimal_and_quantity_integral(self):
        row=normalise(['1','A','Item',3,datetime(2011,1,1),0.1,123.0,'UK'])
        self.assertEqual(Decimal(row[5])*row[3],Decimal('.3'))
        self.assertEqual(row[6],'123')
    def test_fractional_quantity_rejected(self):
        with self.assertRaises(ValueError): normalise(['1','A','x',1.2,datetime(2011,1,1),1,None,'UK'])
    def test_invalid_date_rejected(self):
        with self.assertRaises(ValueError): normalise(['1','A','x',1,None,1,None,'UK'])
    def test_nan_price_rejected(self):
        with self.assertRaises(ValueError): normalise(['1','A','x',1,datetime(2011,1,1),float('nan'),None,'UK'])
    def test_excess_price_precision_rejected(self):
        with self.assertRaises(ValueError): normalise(['1','A','x',1,datetime(2011,1,1),'.12345',None,'UK'])
    def test_duplicate_returns_missingness_and_reload(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td); source=p/'fixture.csv'
            sale=['1','A','Toy',3,'2011-01-01 00:00:00','.1','','UK']
            returned=['C2','A','Toy',-1,'2011-01-02 00:00:00','.1','7','UK']
            zero=['3','B','',2,'2011-01-03 00:00:00','0','7','France']
            with source.open('w',newline='') as f:
                w=csv.writer(f);w.writerow(FIELDS);w.writerows([sale,sale,returned,zero])
            first=build(source,p/'test.duckdb',p/'test.parquet')
            second=build(source,p/'test.duckdb',p/'test.parquet')
            self.assertEqual(first,second)
            self.assertEqual(first['raw_rows'],4)
            self.assertEqual(first['deduplicated_rows'],3)
            self.assertEqual(first['duplicate_rows_removed'],1)
            self.assertEqual(first['negative_quantity_rows'],1)
            self.assertEqual(first['cancellation_rows'],1)
            self.assertEqual(first['missing_customer_rows'],1)
            self.assertEqual(first['missing_description_rows'],1)
            self.assertEqual(first['sale_line_rows'],1)
            self.assertEqual(Decimal(first['raw_signed_value']),Decimal('.5'))
            self.assertEqual(Decimal(first['deduplicated_signed_value']),Decimal('.2'))
            self.assertEqual(Decimal(first['positive_non_cancel_line_value']),Decimal('.3'))
            self.assertEqual(query(p/'test.duckdb','line_count','UK')['value'],'2')
            self.assertEqual(query(p/'test.duckdb','line_count',"UK'; DROP TABLE transactions; --")['value'],'0')
            self.assertEqual(query(p/'test.duckdb','line_count')['value'],'3')
            with self.assertRaises(ValueError): query(p/'test.duckdb','DROP TABLE transactions')
            with self.assertRaises(ValueError): query(p/'test.duckdb','line_count','x'*81)
    def test_both_sheets_are_streamed(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td); wb=openpyxl.Workbook(); ws=wb.active
            for sheet in [ws,wb.create_sheet('second')]:
                sheet.append(['Invoice','StockCode','Description','Quantity','InvoiceDate','Price','Customer ID','Country'])
                sheet.append(['1','A','Toy',2,datetime(2011,1,1),1.25,None,'UK'])
            wb.save(p/'input.xlsx')
            with zipfile.ZipFile(p/'source.zip','w') as z:z.write(p/'input.xlsx','some/path/retail.xlsx')
            r=workbook_to_csv(p/'source.zip',p/'staging'/'data.csv')
            self.assertEqual(r['raw_rows'],2)
            self.assertEqual(Decimal(r['python_decimal_raw_signed_value']),Decimal('5'))
    def test_schema_drift_fails_explicitly(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td); wb=openpyxl.Workbook();wb.active.append(['wrong']);wb.save(p/'input.xlsx')
            with zipfile.ZipFile(p/'source.zip','w') as z:z.write(p/'input.xlsx','input.xlsx')
            with self.assertRaisesRegex(ValueError,'schema'):
                workbook_to_csv(p/'source.zip',p/'staging'/'data.csv')

if __name__=='__main__': unittest.main(verbosity=2)
